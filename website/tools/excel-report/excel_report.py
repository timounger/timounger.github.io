"""!
********************************************************************************
@file   excel_report.py
@brief  Create an Excel print report (PrintReport.xlsx) from a PrintLog.csv,
        modelled on the reference project's excel_print_report.py.

        Sheets: print log, summary, article bar chart, cash report, tool info.
        Run standalone (bundled as excel_report.exe and called by the desktop
        app when a report is created):
            python excel_report.py <print_log.csv> <out.xlsx> <articles.ini>

        Requirements: Python 3.9+, openpyxl.
********************************************************************************
"""

import sys
import csv
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

UNIT = "EUR"
DATE_FORMAT = "%Y/%m/%d %H:%M:%S"
NUMBER_FORMAT_EUR = '#,##0.00 "EUR"'
NUMBER_FORMAT_PERCENT = "0%"
NUMBER_FORMAT_DATETIME = "YYYY-MM-DD HH:MM:SS"
THIN = Side(style="thin")
THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT_NAME = "Consolas"
FONT_SIZE = 11

TAX_GROUPS = (1, 2, 3)
FREE_USER = "Free"
LOCAL_USER = "Local"
# Cash denominations for the counting sheet (EUR).
CASH_VALUES = [500, 200, 100, 50, 20, 10, 5, 2, 1, 0.5, 0.2, 0.1, 0.05, 0.02, 0.01]
GROUP_LABELS = {1: "Getränke", 2: "Speisen", 3: "Sonstiges"}


def set_cell(ws, row, col, value, bold=False, number_format=None, border=None, fill_color=None):
    """!
    @brief Write a styled cell.
    @param ws : worksheet
    @param row : 1-based row
    @param col : 1-based column
    @param value : cell value
    @param bold : bold font
    @param number_format : optional number format string
    @param border : optional border
    @param fill_color : optional solid fill color (hex)
    @return the written cell
    """
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(name=FONT_NAME, size=FONT_SIZE, bold=bold)
    if number_format:
        cell.number_format = number_format
    if border:
        cell.border = border
    if fill_color:
        cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    return cell


def read_section(path, section):
    """!
    @brief Read the key/value entries of a named INI section (e.g. [Tax]).
    @param path : articles.ini path
    @param section : section name without brackets
    @return dict of lower-cased key to value
    """
    entries = {}
    in_section = False
    target = f"[{section.lower()}]"
    try:
        with open(path, mode="r", encoding="utf-8") as ini:
            for raw in ini:
                line = raw.strip()
                if not line or line.startswith(("#", ";")):
                    continue
                if line.startswith("[") and line.endswith("]"):
                    in_section = line.lower() == target
                elif in_section and "=" in line:
                    key, _, value = line.partition("=")
                    entries[key.strip().lower()] = value.strip()
    except OSError:
        pass
    return entries


def read_items(csv_path):
    """!
    @brief Read the print log CSV into item dicts.
    @param csv_path : PrintLog.csv path
    @return list of item dicts (amount, name, pos, group, total, user, date, port)
    """
    items = []
    try:
        with open(csv_path, mode="r", encoding="utf-8", newline="") as csv_file:
            for i, row in enumerate(csv.reader(csv_file, delimiter=";")):
                if i == 0 or len(row) < 8:  # skip header / malformed
                    continue
                try:
                    amount = int(row[0])
                    total = float(row[4])
                    date = datetime.strptime(row[6], DATE_FORMAT)
                except ValueError:
                    continue
                group = int(row[3]) if row[3].isdigit() else 3
                if group not in (1, 2):
                    group = 3
                items.append({"amount": amount, "name": row[1], "pos": row[2], "group": group,
                              "total": total, "user": row[5], "date": date, "port": row[7]})
    except OSError:
        pass
    return items


def create_print_log_sheet(ws, items):
    """! Print-log table sheet. @param ws : worksheet @param items : item dicts """
    ws.title = "Druckprotokoll"
    ws.sheet_properties.tabColor = "00B050"
    ws.freeze_panes = "A2"
    header = ["Anzahl", "Artikelname", "Artikelnummer", "Gruppe", "Preis", "Gesamtpreis", "Benutzer", "Datum", "Drucker"]
    for col, heading in enumerate(header, start=1):
        set_cell(ws, 1, col, heading, bold=True, border=THIN_BORDER)
    for i, item in enumerate(items):
        row = i + 2
        set_cell(ws, row, 1, item["amount"], border=THIN_BORDER)
        set_cell(ws, row, 2, item["name"], border=THIN_BORDER)
        set_cell(ws, row, 3, item["pos"], border=THIN_BORDER)
        set_cell(ws, row, 4, item["group"], border=THIN_BORDER)
        set_cell(ws, row, 5, f"=(F{row}/A{row})", number_format=NUMBER_FORMAT_EUR, border=THIN_BORDER)
        set_cell(ws, row, 6, item["total"], number_format=NUMBER_FORMAT_EUR, border=THIN_BORDER)
        set_cell(ws, row, 7, item["user"], border=THIN_BORDER)
        set_cell(ws, row, 8, item["date"], number_format=NUMBER_FORMAT_DATETIME, border=THIN_BORDER)
        set_cell(ws, row, 9, item["port"], border=THIN_BORDER)
    for col, width in enumerate([6, 18, 12, 7, 10, 12, 10, 22, 8], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def aggregate(items):
    """!
    @brief Aggregate items per user/group and per article (storno nets out).
    @param items : item dicts
    @return (user_items, total_items, free_items, printed_articles, group_sums, min_date, max_date)
    """
    user_items = {}
    total_items = {g: {"sum": 0.0, "sales": {}} for g in TAX_GROUPS}
    free_items = {g: {"sum": 0.0, "sales": {}} for g in TAX_GROUPS}
    group_sums = {g: 0.0 for g in TAX_GROUPS}
    min_date = None
    max_date = None
    for item in items:
        count = item["amount"] if item["total"] >= 0 else -item["amount"]  # storno subtracts
        price = item["total"]
        group = item["group"]
        name = item["name"]
        user = item["user"]
        min_date = item["date"] if min_date is None else min(min_date, item["date"])
        max_date = item["date"] if max_date is None else max(max_date, item["date"])
        group_sums[group] += price
        if user != FREE_USER:
            entry = user_items.setdefault(user, {g: {"sum": 0.0, "sales": {}} for g in TAX_GROUPS})
            entry[group]["sales"][name] = entry[group]["sales"].get(name, 0) + count
            entry[group]["sum"] += price
            total_items[group]["sales"][name] = total_items[group]["sales"].get(name, 0) + count
            total_items[group]["sum"] += price
        else:
            free_items[group]["sales"][name] = free_items[group]["sales"].get(name, 0) + count
    printed_articles = {}
    for group in TAX_GROUPS:
        for name, count in total_items[group]["sales"].items():
            printed_articles[name] = printed_articles.get(name, 0) + count
    return user_items, total_items, free_items, printed_articles, group_sums, min_date, max_date


def group_summary_lines(group_dict, print_values=True):
    """! Format per-group sales lines. @param group_dict : groups @param print_values : show sums @return list of text lines """
    lines = []
    for group in TAX_GROUPS:
        group_sum = group_dict[group]["sum"]
        sales = group_dict[group]["sales"]
        if group_sum != 0 or sales:
            if print_values:
                lines.append(f"* Gruppe {group}: {group_sum:.2f} {UNIT}")
            else:
                lines.append(f"* Gruppe {group}:")
        for name, count in sorted(sales.items()):
            lines.append(f"  {count} x {name}")
    return lines


def build_summary(items, involvement_percent):
    """!
    @brief Build the summary text lines (totals, per user, free articles).
    @param items : item dicts
    @param involvement_percent : user involvement percentage
    @return list of summary text lines
    """
    user_items, total_items, free_items, _articles, _sums, min_date, max_date = aggregate(items)
    lines = ["Bericht", ""]
    lines.append(f"Erstellt: {datetime.now().strftime(DATE_FORMAT)}")
    if min_date and max_date:
        lines.append(f"Von: {min_date.strftime(DATE_FORMAT)}")
        lines.append(f"Bis: {max_date.strftime(DATE_FORMAT)}")
    lines.append("")
    total_sum = sum(total_items[g]["sum"] for g in TAX_GROUPS)
    lines.append("Gesamtverkäufe")
    lines.append(f"Summe: {total_sum:.2f} {UNIT}")
    lines += group_summary_lines(total_items)
    lines.append("")
    lines.append("Verkäufe pro Benutzer")
    for user in sorted(user_items.keys()):
        user_sum = sum(user_items[user][g]["sum"] for g in TAX_GROUPS)
        lines.append(f"# {user}")
        lines.append(f"Summe: {user_sum:.2f} {UNIT}")
        if (not user.startswith(LOCAL_USER)) and involvement_percent != 0:
            involvement = user_sum * (involvement_percent / 100)
            lines.append(f"Beteiligung: {involvement:.2f} {UNIT} ({involvement_percent} %)")
            lines.append(f"Auszahlung: {user_sum - involvement:.2f} {UNIT}")
        lines += group_summary_lines(user_items[user])
    if any(free_items[g]["sales"] for g in TAX_GROUPS):
        lines.append("")
        lines.append("Gratis-Artikel")
        lines += group_summary_lines(free_items, print_values=False)
    return lines


def create_summary_sheet(ws, summary_lines):
    """! Summary text sheet. @param ws : worksheet @param summary_lines : text lines """
    ws.sheet_properties.tabColor = "000000"
    for i, line in enumerate(summary_lines, start=1):
        set_cell(ws, i, 1, line, bold=line and not line.startswith((" ", "*", "#")))
    ws.column_dimensions["A"].width = 50


def create_article_chart_sheet(ws, printed_articles):
    """! Article count bar-chart sheet. @param ws : worksheet @param printed_articles : name to count """
    ws.sheet_properties.tabColor = "FFA500"
    set_cell(ws, 1, 1, "Artikel", bold=True, border=THIN_BORDER)
    set_cell(ws, 1, 2, "Drucke", bold=True, border=THIN_BORDER)
    set_cell(ws, 2, 1, "Summe", bold=True, border=THIN_BORDER)
    row = 2
    for name, count in sorted(printed_articles.items(), key=lambda x: x[1], reverse=True):
        row += 1
        set_cell(ws, row, 1, name, border=THIN_BORDER)
        set_cell(ws, row, 2, count, border=THIN_BORDER)
    set_cell(ws, 2, 2, f"=SUM(B3:B{row})", bold=True, border=THIN_BORDER)
    if row >= 3:
        chart = BarChart()
        chart.type = "col"
        chart.title = "Gedruckte Artikel"
        chart.legend = None
        chart.height = 15
        chart.width = 20
        data = Reference(ws, min_col=2, min_row=3, max_row=row)
        cats = Reference(ws, min_col=1, min_row=3, max_row=row)
        chart.add_data(data, titles_from_data=False)
        chart.set_categories(cats)
        ws.add_chart(chart, "D1")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 8


def create_cash_report_sheet(ws, group_values):
    """!
    @brief Cash counting + group gross/net/tax sheet.
    @param ws : worksheet
    @param group_values : list of [gross, tax_percent] per group
    """
    ws.sheet_properties.tabColor = "FF0000"
    set_cell(ws, 1, 1, "Anfang", bold=True)
    set_cell(ws, 1, 5, "Ende", bold=True)
    for i, heading in enumerate(["Kassenwert", "Anzahl", "Summe"]):
        set_cell(ws, 3, i + 1, heading, bold=True, border=THIN_BORDER)
        set_cell(ws, 3, i + 5, heading, bold=True, border=THIN_BORDER)
    for i, cash in enumerate(CASH_VALUES):
        cash_row = i + 4
        set_cell(ws, cash_row, 1, cash, border=THIN_BORDER, number_format=NUMBER_FORMAT_EUR)
        set_cell(ws, cash_row, 5, cash, border=THIN_BORDER, number_format=NUMBER_FORMAT_EUR)
        set_cell(ws, cash_row, 2, 0, border=THIN_BORDER)
        set_cell(ws, cash_row, 6, 0, border=THIN_BORDER)
        set_cell(ws, cash_row, 3, f"=(A{cash_row}*B{cash_row})", border=THIN_BORDER, number_format=NUMBER_FORMAT_EUR)
        set_cell(ws, cash_row, 7, f"=(E{cash_row}*F{cash_row})", border=THIN_BORDER, number_format=NUMBER_FORMAT_EUR)
    set_cell(ws, 19, 1, "Summe", bold=True, border=THIN_BORDER)
    set_cell(ws, 19, 5, "Summe", bold=True, border=THIN_BORDER)
    set_cell(ws, 19, 3, "=SUM(C4:C18)", bold=True, border=THIN_BORDER, number_format=NUMBER_FORMAT_EUR)
    set_cell(ws, 19, 7, "=SUM(G4:G18)", bold=True, border=THIN_BORDER, number_format=NUMBER_FORMAT_EUR)
    set_cell(ws, 21, 1, "Bargeld", border=THIN_BORDER)
    set_cell(ws, 21, 2, "=(G19-C19)", border=THIN_BORDER, number_format=NUMBER_FORMAT_EUR)
    set_cell(ws, 22, 1, "Gezählt Druck", border=THIN_BORDER)
    set_cell(ws, 22, 2, "=SUM(B26:B28)", border=THIN_BORDER, number_format=NUMBER_FORMAT_EUR)
    set_cell(ws, 23, 1, "Differenz", border=THIN_BORDER)
    set_cell(ws, 23, 2, "=IF(OR(C19<>0,G19<>0),(B21-B22),0)", border=THIN_BORDER, number_format=NUMBER_FORMAT_EUR)
    for i, heading in enumerate([None, "Brutto", "Netto", "Steuersatz", "Steuer"]):
        set_cell(ws, 25, i + 1, heading, bold=True, border=THIN_BORDER)
    for i, values in enumerate(group_values):
        report_row = 26 + i
        set_cell(ws, report_row, 1, f"Gruppe {i + 1} {GROUP_LABELS.get(i + 1, '')}", border=THIN_BORDER)
        set_cell(ws, report_row, 2, values[0], border=THIN_BORDER, number_format=NUMBER_FORMAT_EUR)
        set_cell(ws, report_row, 3, f"=(B{report_row}/(1+D{report_row}))", border=THIN_BORDER, number_format=NUMBER_FORMAT_EUR)
        set_cell(ws, report_row, 4, values[1] / 100, border=THIN_BORDER, number_format=NUMBER_FORMAT_PERCENT)
        set_cell(ws, report_row, 5, f"=(B{report_row}-C{report_row})", border=THIN_BORDER, number_format=NUMBER_FORMAT_EUR)
    set_cell(ws, 30, 1, "Summe", bold=True, border=THIN_BORDER)
    set_cell(ws, 30, 2, "=SUM(B26:B29)", bold=True, border=THIN_BORDER, number_format=NUMBER_FORMAT_EUR, fill_color="92D050")
    set_cell(ws, 30, 5, "=SUM(E26:E29)", bold=True, border=THIN_BORDER, number_format=NUMBER_FORMAT_EUR)
    for col, width in [("A", 20), ("B", 16), ("C", 16), ("D", 14), ("E", 20), ("F", 16), ("G", 16)]:
        ws.column_dimensions[col].width = width


def create_tool_info_sheet(ws):
    """! Hidden tool-info sheet. @param ws : worksheet """
    ws.sheet_state = "hidden"
    set_cell(ws, 1, 1, f"Dieser Bericht wurde erstellt am {datetime.now().strftime(DATE_FORMAT)} mit:")
    set_cell(ws, 3, 1, "BonPrinter", bold=True)
    set_cell(ws, 4, 1, "BonPrinter - Thermal printer tool")
    ws.column_dimensions["A"].width = 60


def main():
    """! Read the CSV + config and write the Excel report. """
    if len(sys.argv) < 3:
        print("usage: excel_report.py <print_log.csv> <out.xlsx> [articles.ini]")
        return
    csv_path = sys.argv[1]
    out_path = sys.argv[2]
    ini_path = sys.argv[3] if len(sys.argv) > 3 else None

    items = read_items(csv_path)
    tax = read_section(ini_path, "Tax") if ini_path else {}

    def tax_percent(key):
        """! Parse a tax percentage from the [Tax] section. @param key : key name @return percent """
        try:
            return float(tax.get(key, "0"))
        except ValueError:
            return 0.0

    involvement = tax_percent("user")
    _u, _t, _f, printed_articles, group_sums, _min, _max = aggregate(items)
    group_values = [[group_sums[g], tax_percent(f"group{g}")] for g in TAX_GROUPS]

    workbook = Workbook()
    create_print_log_sheet(workbook.active, items)
    create_summary_sheet(workbook.create_sheet("Übersicht"), build_summary(items, involvement))
    if printed_articles:
        create_article_chart_sheet(workbook.create_sheet("Artikel"), printed_articles)
        create_cash_report_sheet(workbook.create_sheet("Kassenbericht"), group_values)
    create_tool_info_sheet(workbook.create_sheet("ToolInfo"))
    workbook.save(out_path)
    print(f"Excel report written: {out_path}")


if __name__ == "__main__":
    main()
